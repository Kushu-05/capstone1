import os
import pytest
from pathlib import Path
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend.config import settings
from backend.database import Base, get_db
from backend.main import app

# In-memory test SQLite DB
TEST_DATABASE_URL = "sqlite:///:memory:"
engine_test = create_engine(TEST_DATABASE_URL, connect_args={"check_same_thread": False})
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine_test)

@pytest.fixture(scope="session", autouse=True)
def init_test_db():
    Base.metadata.create_all(bind=engine_test)
    yield
    Base.metadata.drop_all(bind=engine_test)

@pytest.fixture
def db_session():
    connection = engine_test.connect()
    transaction = connection.begin()
    session = TestingSessionLocal(bind=connection)
    try:
        yield session
    finally:
        session.close()
        transaction.rollback()
        connection.close()

@pytest.fixture
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()

@pytest.fixture
def sample_txt_file(tmp_path: Path) -> Path:
    p = tmp_path / "sample_report.txt"
    p.write_text(
        "SECTION 1: OVERVIEW\n"
        "The committee reviewed $15 million in research expenditures across 12 facilities.\n\n"
        "SECTION 2: RISKS\n"
        "A critical compliance risk was identified regarding missing milestone reports.\n\n"
        "SECTION 3: RECOMMENDATIONS\n"
        "We recommend mandatory quarterly audits for all programs.",
        encoding="utf-8"
    )
    return p

@pytest.fixture
def sample_pdf_file(tmp_path: Path) -> Path:
    import fitz
    p = tmp_path / "sample_audit.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 72),
        "SECTION 1: EXECUTIVE AUDIT SUMMARY\n\n"
        "The federal grant initiative allocated $28 million to regional technology labs.\n"
        "Our review found a severe risk of supply chain delays affecting 42% of deliverables.\n\n"
        "SECTION 2: DECISIONS & RECOMMENDATIONS\n\n"
        "The director decided to halt non-compliant contracts by December 31, 2023.\n"
        "We recommend deploying automated progress dashboards."
    )
    doc.save(str(p))
    doc.close()
    return p

@pytest.fixture
def sample_docx_file(tmp_path: Path) -> Path:
    import docx
    p = tmp_path / "sample_briefing.docx"
    doc = docx.Document()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("This briefing assesses emerging cyber defense protocols across critical sectors.")
    doc.add_heading("Risks and Vulnerabilities", level=2)
    doc.add_paragraph("Legacy operational systems show critical vulnerabilities to ransomware attacks.")
    doc.add_heading("Action Items", level=2)
    doc.add_paragraph("Deploy security patches across all endpoints by November 15, 2023.")
    doc.save(str(p))
    return p
